"""Create read-only-derived Adobe Document Generation templates for parking and warehouse leases.

The originals are only opened for reading.  Each generated template retains the
approved source layout and legal language, while contract fields and the legal
terms body are replaced with Adobe tags.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph


def replace_paragraph(paragraph: Paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def replace_cell(cell: _Cell, value: str) -> None:
    replace_paragraph(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        replace_paragraph(paragraph, "")


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def replace_terms_with_tag(document: Document) -> str:
    paragraphs = list(document.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip().startswith("第１条（"))
    end = next(
        i
        for i, p in enumerate(paragraphs[start + 1 :], start + 1)
        if p.text.strip().startswith("以下余白") or p.text.strip().startswith("本契約の証として")
    )
    terms = "\n".join(p.text.strip() for p in paragraphs[start:end] if p.text.strip())
    replace_paragraph(paragraphs[start], "{{termsText}}")
    for paragraph in paragraphs[start + 1 : end]:
        remove_paragraph(paragraph)
    return terms


def fill_summary(document: Document, values: dict[str, str]) -> None:
    summary = document.tables[0]
    for row in summary.rows:
        label = row.cells[0].text.replace("\n", "")
        for expected, value in values.items():
            if expected in label:
                replace_cell(row.cells[1], value)
                break


def replace_signature_slots(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text
        stripped = text.strip()
        if stripped == "賃借人" or stripped == "賃借人：":
            replace_paragraph(paragraph, text.replace("賃借人", "賃借人　{{tenantName}}", 1))
        elif "賃借人：" in text and "{{tenantName}}" not in text:
            replace_paragraph(paragraph, text.replace("賃借人：", "賃借人：　{{tenantName}}", 1))
        elif "賃貸人 ＳＫハウジング株式会社と、賃借人 ○○との間に" in text:
            replace_paragraph(paragraph, text.replace("賃借人 ○○", "賃借人 {{tenantName}}"))


def create_template(source: Path, destination: Path, kind: str) -> str:
    document = Document(source)
    if not document.tables:
        raise ValueError("The approved source must have a contract-summary table.")

    common = {
        "賃借人": "{{tenantName}}",
        "建物名称": "{{propertyName}}",
        "所在　　（地番）": "{{propertyLotAddress}}",
        "（住居表示）": "{{propertyAddress}}",
        "賃貸借期間": "{{contractStartDate}} から {{contractEndDate}} まで\n（{{contractDurationYears}}年）",
        "賃料の支払期日": "{{rentPaymentDue}}",
        "日割計算方法": "{{dailyCalculationMethod}}",
        "敷金": "金 {{depositAmount}} 円（{{depositMonths}}ヶ月分）",
        "特約事項": "{{specialProvisions}}",
    }
    if kind == "parking":
        values = {
            **common,
            "使用目的": "{{usePurpose}}",
            "契約台数": "{{parkingContractCount}} 台",
            "賃料（": "月額金 {{monthlyRentAmount}} 円（税別）\n（1台あたり金 {{monthlyRentPerSpaceAmount}} 円（税別））",
            "車庫証明発行料": "1通あたり金 {{parkingCertificateFeeAmount}} 円（税別）",
        }
    elif kind == "warehouse":
        values = {
            **common,
            "連帯保証人": "{{guarantorName}}（極度額 {{guarantorLimitAmount}} 円を上限とする）",
            "構造規模": "{{buildingStructure}}",
            "賃貸借物件": "{{floorLabel}}階　{{unitNames}}\n{{leasedAreaSqm}}㎡（{{leasedAreaTsubo}}坪）",
            "使用目的": "{{usePurpose}}",
            "賃料（": "月額金 {{monthlyRentAmount}} 円（税別）",
        }
    else:
        raise ValueError(f"Unsupported kind: {kind}")

    fill_summary(document, values)
    replace_signature_slots(document)
    terms = replace_terms_with_tag(document)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return terms


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("--kind", choices=("parking", "warehouse"), required=True)
    parser.add_argument("--terms-output", type=Path, required=True)
    args = parser.parse_args()
    terms = create_template(args.source, args.destination, args.kind)
    args.terms_output.parent.mkdir(parents=True, exist_ok=True)
    args.terms_output.write_text(terms, encoding="utf-8")


if __name__ == "__main__":
    main()
