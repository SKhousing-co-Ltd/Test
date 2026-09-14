"""Extract the approved fixed-term lease source into reviewable Markdown."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    document = Document(args.source)
    lines = ["# 定期建物賃貸借契約書 原文抽出", "", f"抽出元: `{args.source}`", ""]
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            lines.extend([f"<!-- paragraph:{index} -->", text, ""])
    for table_index, table in enumerate(document.tables, start=1):
        lines.extend([f"## 表 {table_index}", ""])
        for row in table.rows:
            lines.append(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
        lines.append("")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
